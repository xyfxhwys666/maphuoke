from __future__ import annotations

import csv
from dataclasses import dataclass
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any, ClassVar

from backend.app.domain.dedup import company_identity
from backend.app.domain.normalization import normalize_company_name
from openpyxl import load_workbook


class ImportValidationError(ValueError):
    pass


@dataclass(frozen=True, slots=True)
class CompanyInput:
    company_name: str
    city: str
    row_number: int


class CompanyInputImporter:
    _COMPANY_HEADERS: ClassVar[set[str]] = {
        'company_name',
        'company',
        'name',
        '公司名称',
        '企业名称',
        '公司名',
    }
    _CITY_HEADERS: ClassVar[set[str]] = {'city', '城市', '所在城市', '城市名称'}

    @classmethod
    def from_csv(cls, source: str | bytes | Path) -> list[CompanyInput]:
        text = cls._csv_text(source)
        reader = csv.reader(StringIO(text))
        try:
            header_row = next(reader)
        except StopIteration as exc:
            raise ImportValidationError('CSV 文件为空') from exc
        indexes = cls._header_indexes(header_row)
        return cls._rows_from_values(reader, indexes)

    @classmethod
    def from_xlsx(cls, source: bytes | bytearray | BytesIO | Path) -> list[CompanyInput]:
        if isinstance(source, Path):
            workbook_source: Any = source
        else:
            workbook_source = BytesIO(bytes(source) if not hasattr(source, 'read') else source.read())
        workbook = load_workbook(workbook_source, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            rows = sheet.iter_rows(values_only=True)
            header_row = [str(value or '') for value in next(rows)]
            indexes = cls._header_indexes(header_row)
            return cls._rows_from_values(rows, indexes)
        except StopIteration as exc:
            raise ImportValidationError('XLSX 文件为空') from exc
        finally:
            workbook.close()

    @staticmethod
    def _csv_text(source: str | bytes | Path) -> str:
        if isinstance(source, Path) or (isinstance(source, str) and '\n' not in source and Path(source).exists()):
            return Path(source).read_text(encoding='utf-8-sig')
        if isinstance(source, bytes):
            return source.decode('utf-8-sig')
        return str(source)

    @classmethod
    def _header_indexes(cls, header_row: list[str]) -> tuple[int, int]:
        normalized = {str(value).strip().casefold(): index for index, value in enumerate(header_row)}
        company_index = next((normalized[key] for key in cls._COMPANY_HEADERS if key.casefold() in normalized), None)
        city_index = next((normalized[key] for key in cls._CITY_HEADERS if key.casefold() in normalized), None)
        if company_index is None or city_index is None:
            raise ImportValidationError('缺少必填列：公司名称、城市')
        return company_index, city_index

    @staticmethod
    def _rows_from_values(rows, indexes: tuple[int, int]) -> list[CompanyInput]:
        company_index, city_index = indexes
        result: list[CompanyInput] = []
        seen: set[str] = set()
        for row_number, row in enumerate(rows, start=2):
            values = list(row)
            company = normalize_company_name(values[company_index] if company_index < len(values) else None) or ''
            city = str(values[city_index] if city_index < len(values) and values[city_index] is not None else '').strip()
            if not company or not city:
                raise ImportValidationError(f'第{row_number}行缺少公司名称或城市')
            identity = company_identity(city, company)
            if identity in seen:
                continue
            seen.add(identity)
            result.append(CompanyInput(company_name=company, city=city, row_number=row_number))
        return result
